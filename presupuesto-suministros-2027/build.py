"""Genera el dashboard de defensa del Presupuesto de Mantención Suministros 2027.

Uso:  python3 build.py <ruta/Presupuesto_2027-Afta.xlsm>

Lee el Excel (valores calculados), limpia y normaliza, y escribe:
  - index.html              dashboard único (datos embebidos)
  - trazabilidad.csv        área, prioridad, actividad, monto, fila origen
  - datos_pendientes.md     insumos faltantes / inconsistencias detectadas
Ningún monto se estima: todo proviene de una celda del Excel.
"""
import csv, json, re, sys, unicodedata, warnings
from pathlib import Path
import openpyxl

warnings.filterwarnings("ignore")
HERE = Path(__file__).parent
AREAS = ["Planta de Agua", "Planta de RILES", "Sala Compresores", "Planta de Frío", "SSEE", "Sala de Caldera"]


def norm_key(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


AREA_MAP = {
    "planta de agua": "Planta de Agua",
    "planta de riles": "Planta de RILES",
    "sala compresores": "Sala Compresores",
    "planta de frio": "Planta de Frío",
    "ssee": "SSEE",
    "sala de caldera": "Sala de Caldera",
    "caldera": "Sala de Caldera",
}


def num(v):
    if isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    return 0.0  # None, '$ -', '#REF!' y otros textos se ignoran


def clean(v):
    if v is None:
        return ""
    s = re.sub(r"\s+", " ", str(v)).strip()
    return "" if s.startswith("#") else s


def main(xlsm):
    wb = openpyxl.load_workbook(xlsm, data_only=True)
    pend, notes = [], []

    # ---------- SUM 2027 ----------
    ws = wb["SUM 2027"]
    header_prio = clean(ws.cell(1, 1).value)
    acts, last_area, excluded_total_rows = [], None, []
    for i, r in enumerate(ws.iter_rows(min_row=2, max_col=22, values_only=True), start=2):
        meses = [num(x) for x in r[9:21]]
        total = num(r[21])
        if total == 0:
            continue
        raw_area = clean(r[2])
        if not any(clean(x) for x in r[:9]):  # fila de totales sin atributos
            excluded_total_rows.append((i, total, sum(meses)))
            continue
        flag = ""
        if raw_area:
            area = AREA_MAP.get(norm_key(raw_area))
            if area is None:
                raise SystemExit(f"Área no reconocida en fila {i}: {raw_area!r}")
        else:
            area = last_area
            flag = "Área vacía en Excel; asignada por fila contigua (bloque " + (last_area or "?") + ")"
        last_area = area
        p = r[0]
        prio = {1: "P1", 2: "P2", 3: "P3"}.get(int(p), "SC") if isinstance(p, (int, float)) else "SC"
        if abs(sum(meses) - total) > 1:
            flag = (flag + "; " if flag else "") + f"TOTAL ≠ suma meses ({sum(meses):,.0f})"
        acts.append(dict(
            fila=i, area=area, prio=prio, cc=clean(r[1]), cuenta=clean(r[3]).replace(".0", ""),
            nombreCuenta=clean(r[4]), equipo=clean(r[5]), componente=clean(r[6]), tipo=clean(r[7]),
            desc=clean(r[8]), meses=[round(x) for x in meses], total=round(total), flag=flag))

    for i, t, s in excluded_total_rows:
        notes.append(f"Fila {i} de 'SUM 2027' es una fila de totales (TOTAL {t:,.0f}; suma de meses {s:,.0f}); "
                     "se excluye para no duplicar. El dashboard suma las filas de detalle.")
    det_total = sum(a["total"] for a in acts)
    for i, t, s in excluded_total_rows:
        if abs(t - s) > 1:
            pend.append(f"Fila {i} 'SUM 2027': la columna TOTAL ({t:,.0f}) no cuadra con la suma de Ene–Dic de esa misma fila "
                        f"({s:,.0f}); diferencia {s - t:,.0f} CLP. El detalle suma {det_total:,.0f}. Revisar fórmula de junio.")
    inferred = [a for a in acts if a["flag"].startswith("Área vacía")]
    if inferred:
        pend.append("Filas " + ", ".join(str(a["fila"]) for a in inferred) +
                    f" de 'SUM 2027' no tienen Área ni CC ({sum(a['total'] for a in inferred):,.0f} CLP). Se muestran en "
                    f"{inferred[0]['area']} por contigüidad y quedan marcadas; confirmar.")
    sc = [a for a in acts if a["prio"] == "SC"]
    pend.append(f"{len(sc)} actividades ({sum(a['total'] for a in sc)/1e6:,.1f} MM CLP, "
                f"{sum(a['total'] for a in sc)/det_total:.0%} del total) no tienen prioridad en 'SUM 2027'. "
                "Se muestran como 'Sin clasificar'; no se infiere prioridad.")

    # ---------- Presupuesto 2026 ----------
    ws = wb["Presupuesto 2026"]
    p26 = []
    for i, r in enumerate(ws.iter_rows(min_row=2, max_col=20, values_only=True), start=2):
        total = num(r[19])
        if total == 0 or not clean(r[1]):
            continue
        area = AREA_MAP.get(norm_key(r[1]))
        if area is None:
            continue  # otras áreas de planta (líneas, administración, etc.)
        p26.append(dict(fila=i, area=area, equipo=clean(r[4]), tipo=clean(r[5]), desc=clean(r[6]),
                        meses=[round(num(x)) for x in r[7:19]], total=round(total)))

    # ---------- Proyección de Gastos ----------
    ws = wb["Proyección de Gastos"]
    g = lambda ref: ws[ref].value
    proy = dict(
        mant_real=g("B6"), mant_ppto27=g("C6"), mant_vol27=g("E6"), mant_dif=g("F6"),
        prod26_real=g("B20"), prod26_ppto=g("C20"), prod27=g("F20"),
        sum_real=g("B29"), sum_ppto=g("C29"), sum_pond=g("E29"), mant_total_c37=g("C37"),
    )
    pend.append(f"'Proyección de Gastos' asigna a Suministros {proy['sum_ppto']/1e6:,.1f} MM (C29) dentro de Mantención "
                f"{proy['mant_ppto27']/1e6:,.1f} MM, pero el detalle 'SUM 2027' suma {det_total/1e6:,.1f} MM. "
                "Aclarar cuál es la cifra oficial a defender y el periodo de la columna 'Real' (B29).")

    # ---------- Datos pendientes estructurales ----------
    pend += [
        "N° de motobombas físicas y cuántas operan sin respaldo +1 (referencia verbal ~45 / ~25; no está en el Excel).",
        "Configuración de redundancia (N / N+1) por equipo crítico — el Excel sólo la explicita en algunos equipos de RILES.",
        "Descomposición del puente 2026→2027 en efecto volumen, nuevas actividades, reprogramación y precio/IPC "
        "(el Excel no trae el vínculo actividad 2026 ↔ 2027 ni el supuesto de IPC).",
        "Capacidad de HH disponibles por mes (operadores mantenedores + contratistas) para el gráfico de carga vs capacidad.",
        "Número y año de la norma aplicable a caldera / línea GNL, antes de citarla ante gerencia.",
        "Gasto real 2026 por área de Suministros (sólo existe el presupuesto 2026 por área).",
    ]

    data = dict(
        generado=str(Path(xlsm).name), areas=AREAS, criterioPrioridad=header_prio,
        acts=acts, p26=p26, proy=proy, pendientes=pend, notas=notes,
    )

    tpl = (HERE / "template.html").read_text(encoding="utf-8")
    out = tpl.replace("/*__DATA__*/null", json.dumps(data, ensure_ascii=False))
    (HERE / "index.html").write_text(out, encoding="utf-8")

    with open(HERE / "trazabilidad.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Área", "Prioridad", "Equipo", "Componente", "Actividad", "Monto 2027 (CLP)", "Meses con gasto",
                    "Hoja", "Fila origen", "Observación"])
        mes = "Ene Feb Mar Abr May Jun Jul Ago Sep Oct Nov Dic".split()
        for a in sorted(acts, key=lambda a: (AREAS.index(a["area"]), a["prio"], -a["total"])):
            w.writerow([a["area"], a["prio"] if a["prio"] != "SC" else "Sin clasificar", a["equipo"], a["componente"],
                        a["desc"], a["total"], " ".join(m for m, v in zip(mes, a["meses"]) if v), "SUM 2027", a["fila"], a["flag"]])
        for a in p26:
            w.writerow([a["area"], "(2026, sin prioridad)", a["equipo"], "", a["desc"], a["total"],
                        " ".join(m for m, v in zip(mes, a["meses"]) if v), "Presupuesto 2026", a["fila"], "Monto 2026"])

    md = ["# Datos pendientes e inconsistencias detectadas", "",
          f"Fuente: `{Path(xlsm).name}` — total detalle SUM 2027: {det_total:,.0f} CLP ({len(acts)} actividades).", ""]
    md += [f"{k}. {p}" for k, p in enumerate(pend, 1)]
    md += ["", "## Notas de limpieza", ""] + [f"- {n}" for n in notes] + [
        "- Nombres de área normalizados (trim, mayúsculas, tildes): p. ej. 'planta de frio ' → 'Planta de Frío'; 'caldera ' (2026) → 'Sala de Caldera'.",
        "- Filas con TOTAL 0 excluidas; textos como '$ -' o '#REF!' en meses se tratan como 0.",
        "- Filas sin prioridad se conservan como 'Sin clasificar'."]
    (HERE / "datos_pendientes.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print(f"OK: {len(acts)} actividades 2027 = {det_total:,.0f} CLP; {len(p26)} filas 2026 = {sum(a['total'] for a in p26):,.0f} CLP")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "Presupuesto_2027-Afta.xlsm")
